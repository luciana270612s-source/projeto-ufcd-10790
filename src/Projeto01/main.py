import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import load_workbook
from datetime import datetime
import os

ARQUIVO = "estoque_limpeza.xlsx"
ABA_ESTOQUE = "Estoque"
ABA_MOV = "Movimentos"

CORES = {
    "fundo": "#F4F8F7", "topo": "#1B5E20", "botao": "#2E7D32",
    "hover": "#1B5E20", "branco": "#FFFFFF", "texto": "#1B1B1B",
    "ok_bg": "#DFF5E1", "ok_fg": "#1B5E20",
    "alerta_bg": "#FFD6D6", "alerta_fg": "#B00020"
}

COLS = ["Código", "Produto", "Categoria", "Unidade", "Estoque Atual", "Estoque Mínimo", "Comprar?"]


def num(v):
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace("€", "").replace(",", ".").strip())
    except ValueError:
        return 0


def fmt(v):
    v = num(v)
    return str(int(v)) if v == int(v) else str(v)


def erro(msg):
    messagebox.showerror("Erro", msg)


def abrir_excel():
    if not os.path.exists(ARQUIVO):
        erro(f"O ficheiro {ARQUIVO} não foi encontrado.\nColoca-o na mesma pasta do programa.")
        return None
    try:
        wb = load_workbook(ARQUIVO)
    except PermissionError:
        erro("Fecha o Excel antes de usar ou guardar alterações.")
        return None
    if ABA_ESTOQUE not in wb.sheetnames or ABA_MOV not in wb.sheetnames:
        erro("O Excel precisa ter as abas 'Estoque' e 'Movimentos'.")
        wb.close()
        return None
    return wb


def guardar(wb):
    try:
        wb.save(ARQUIVO)
        return True
    except PermissionError:
        erro("Fecha o Excel antes de guardar alterações.")
        return False
    finally:
        wb.close()


def linha_produto(ws, codigo):
    codigo = str(codigo).strip()
    for lin in range(2, ws.max_row + 1):
        if str(ws.cell(lin, 1).value).strip() == codigo:
            return lin
    return None


def recalcular(wb, codigo):
    est, mov = wb[ABA_ESTOQUE], wb[ABA_MOV]
    lin = linha_produto(est, codigo)
    if not lin:
        return

    inicial, minimo = num(est.cell(lin, 5).value), num(est.cell(lin, 9).value)
    entradas = saidas = 0
    ultima_compra = ""

    for r in range(2, mov.max_row + 1):
        if str(mov.cell(r, 2).value).strip() != str(codigo).strip():
            continue
        tipo = str(mov.cell(r, 4).value).strip().upper()
        qtd = num(mov.cell(r, 5).value)
        if tipo in ["COMPRA", "AJUSTE +", "ENTRADA"]:
            entradas += qtd
            if tipo == "COMPRA":
                ultima_compra = mov.cell(r, 1).value
        elif tipo in ["RETIRADA", "AJUSTE -", "SAÍDA", "SAIDA"]:
            saidas += qtd

    atual = inicial + entradas - saidas
    est.cell(lin, 6).value = entradas
    est.cell(lin, 7).value = saidas
    est.cell(lin, 8).value = atual
    est.cell(lin, 10).value = "COMPRAR" if atual <= minimo else "OK"
    est.cell(lin, 11).value = ultima_compra


def recalcular_tudo():
    wb = abrir_excel()
    if not wb:
        return False
    ws = wb[ABA_ESTOQUE]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value and ws.cell(r, 2).value:
            recalcular(wb, ws.cell(r, 1).value)
    return guardar(wb)


def produtos():
    wb = abrir_excel()
    if not wb:
        return []
    ws, dados = wb[ABA_ESTOQUE], []
    for r in range(2, ws.max_row + 1):
        if not ws.cell(r, 1).value or not ws.cell(r, 2).value:
            continue
        dados.append({
            "codigo": ws.cell(r, 1).value, "produto": ws.cell(r, 2).value,
            "categoria": ws.cell(r, 3).value or "", "unidade": ws.cell(r, 4).value or "",
            "atual": ws.cell(r, 8).value or 0, "minimo": ws.cell(r, 9).value or 0,
            "status": ws.cell(r, 10).value or "", "obs": ws.cell(r, 12).value or ""
        })
    wb.close()
    return dados


def atualizar(lista=None):
    tabela.delete(*tabela.get_children())
    if lista is None:
        recalcular_tudo()
        lista = produtos()
    for p in lista:
        status = str(p["status"]).upper().strip()
        tabela.insert("", tk.END, values=(
            p["codigo"], p["produto"], p["categoria"], p["unidade"],
            fmt(p["atual"]), fmt(p["minimo"]), status
        ), tags=("alerta" if status == "COMPRAR" else "ok",))


def pesquisar():
    termo = busca.get().lower().strip()
    atualizar([p for p in produtos() if termo in str(p["codigo"]).lower()
               or termo in p["produto"].lower() or termo in p["categoria"].lower()])


def para_comprar():
    atualizar([p for p in produtos() if str(p["status"]).upper().strip() == "COMPRAR"])


def selecionado():
    item = tabela.selection()
    if not item:
        messagebox.showwarning("Aviso", "Seleciona um produto na tabela.")
        return None
    return tabela.item(item[0], "values")


def botao(local, texto, cmd, largura=18):
    return tk.Button(local, text=texto, command=cmd, width=largura, bg=CORES["botao"],
                     fg=CORES["branco"], activebackground=CORES["hover"],
                     activeforeground=CORES["branco"], font=("Arial", 10, "bold"),
                     relief=tk.FLAT, cursor="hand2", padx=5, pady=5)


def nova_janela(titulo, tamanho):
    j = tk.Toplevel(root)
    j.title(titulo)
    j.geometry(tamanho)
    j.configure(bg=CORES["fundo"])
    return j


def campo(janela, texto, padrao=""):
    tk.Label(janela, text=texto, bg=CORES["fundo"], fg=CORES["texto"]).pack(pady=(8, 0))
    e = tk.Entry(janela, width=42)
    e.insert(0, padrao)
    e.pack(pady=3)
    return e


def registrar(tipo):
    prod = selecionado()
    if not prod:
        return
    codigo, nome = prod[0], prod[1]
    j = nova_janela(f"Registrar {tipo}", "400x330")

    tk.Label(j, text=f"Produto: {nome}", bg=CORES["fundo"], fg=CORES["texto"],
             font=("Arial", 11, "bold")).pack(pady=8)
    tk.Label(j, text=f"Código: {codigo}", bg=CORES["fundo"], fg=CORES["texto"]).pack()

    qtd = campo(j, "Quantidade:")
    resp = campo(j, "Responsável:", "Rafael")
    preco = campo(j, "Preço unitário, apenas para compra:")

    def salvar():
        try:
            quantidade = float(qtd.get().replace(",", "."))
            if quantidade <= 0:
                raise ValueError
        except ValueError:
            erro("A quantidade deve ser um número maior que zero.")
            return

        preco_unit = ""
        if tipo == "Compra" and preco.get().strip():
            try:
                preco_unit = float(preco.get().replace(",", "."))
            except ValueError:
                erro("O preço deve ser um número.")
                return

        wb = abrir_excel()
        if not wb:
            return
        ws = wb[ABA_MOV]
        r = ws.max_row + 1
        valores = [datetime.now(), codigo, nome, tipo, quantidade,
                   resp.get().strip() or "Não informado", preco_unit,
                   quantidade * preco_unit if preco_unit != "" else "",
                   "Movimento registrado pelo sistema", f"{codigo}|{tipo}"]
        for c, v in enumerate(valores, 1):
            ws.cell(r, c).value = v
        recalcular(wb, codigo)
        if guardar(wb):
            messagebox.showinfo("Sucesso", f"{tipo} registrada com sucesso.")
            j.destroy()
            atualizar()

    botao(j, "Salvar", salvar, 20).pack(pady=18)


def adicionar():
    j = nova_janela("Adicionar produto", "430x470")
    tk.Label(j, text="Adicionar novo produto", bg=CORES["topo"], fg=CORES["branco"],
             font=("Arial", 13, "bold"), pady=8).pack(fill=tk.X)

    nomes = ["Código", "Produto", "Categoria", "Unidade", "Estoque Inicial", "Estoque Mínimo", "Observações"]
    campos = {n: campo(j, n) for n in nomes}

    def salvar():
        codigo = campos["Código"].get().strip()
        nome = campos["Produto"].get().strip()
        if not codigo or not nome:
            erro("Código e produto são obrigatórios.")
            return
        try:
            inicial = float(campos["Estoque Inicial"].get().replace(",", "."))
            minimo = float(campos["Estoque Mínimo"].get().replace(",", "."))
        except ValueError:
            erro("Estoque inicial e estoque mínimo devem ser números.")
            return

        wb = abrir_excel()
        if not wb:
            return
        ws = wb[ABA_ESTOQUE]
        if linha_produto(ws, codigo):
            erro("Já existe um produto com esse código.")
            wb.close()
            return

        r = ws.max_row + 1
        valores = [codigo, nome, campos["Categoria"].get().strip(), campos["Unidade"].get().strip(),
                   inicial, 0, 0, inicial, minimo, "COMPRAR" if inicial <= minimo else "OK",
                   "", campos["Observações"].get().strip()]
        for c, v in enumerate(valores, 1):
            ws.cell(r, c).value = v
        if guardar(wb):
            messagebox.showinfo("Sucesso", "Produto adicionado com sucesso.")
            j.destroy()
            atualizar()

    botao(j, "Salvar produto", salvar, 22).pack(pady=15)


root = tk.Tk()
root.title("Sistema de Gestão de Stock de Produtos de Limpeza")
root.geometry("1050x600")
root.configure(bg=CORES["fundo"])

style = ttk.Style()
style.theme_use("clam")
style.configure("Treeview", background=CORES["branco"], foreground=CORES["texto"],
                rowheight=30, fieldbackground=CORES["branco"], font=("Arial", 10))
style.configure("Treeview.Heading", background=CORES["topo"], foreground=CORES["branco"],
                font=("Arial", 10, "bold"))
style.map("Treeview", background=[("selected", "#B2DFDB")], foreground=[("selected", CORES["texto"])])

tk.Label(root, text="Sistema de Gestão de Stock de Produtos de Limpeza", font=("Arial", 18, "bold"),
         bg=CORES["topo"], fg=CORES["branco"], pady=14).pack(fill=tk.X)
tk.Label(root, text="Linhas vermelhas = COMPRAR | Linhas verdes = OK", font=("Arial", 10, "bold"),
         bg=CORES["fundo"], fg=CORES["texto"]).pack(pady=8)

frame_busca = tk.Frame(root, bg=CORES["fundo"])
frame_busca.pack(pady=5)
tk.Label(frame_busca, text="Pesquisar produto:", bg=CORES["fundo"], fg=CORES["texto"],
         font=("Arial", 10, "bold")).pack(side=tk.LEFT, padx=5)
busca = tk.Entry(frame_busca, width=42)
busca.pack(side=tk.LEFT, padx=5)
botao(frame_busca, "Pesquisar", pesquisar).pack(side=tk.LEFT, padx=5)
botao(frame_busca, "Mostrar todos", lambda: (busca.delete(0, tk.END), atualizar())).pack(side=tk.LEFT, padx=5)
botao(frame_busca, "Produtos para comprar", para_comprar, 22).pack(side=tk.LEFT, padx=5)

frame_tab = tk.Frame(root, bg=CORES["fundo"])
frame_tab.pack(pady=10, fill=tk.BOTH, expand=True, padx=15)
tabela = ttk.Treeview(frame_tab, columns=COLS, show="headings", height=15)
for col in COLS:
    tabela.heading(col, text=col)

tamanhos = [100, 250, 160, 100, 120, 120, 120]
for col, tam in zip(COLS, tamanhos):
    tabela.column(col, width=tam, anchor=tk.CENTER if col != "Produto" and col != "Categoria" else tk.W)

tabela.tag_configure("alerta", background=CORES["alerta_bg"], foreground=CORES["alerta_fg"])
tabela.tag_configure("ok", background=CORES["ok_bg"], foreground=CORES["ok_fg"])
scroll = ttk.Scrollbar(frame_tab, orient=tk.VERTICAL, command=tabela.yview)
tabela.configure(yscrollcommand=scroll.set)
tabela.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
scroll.pack(side=tk.RIGHT, fill=tk.Y)

frame_btn = tk.Frame(root, bg=CORES["fundo"])
frame_btn.pack(pady=12)
botao(frame_btn, "Adicionar produto", adicionar, 20).pack(side=tk.LEFT, padx=6)
botao(frame_btn, "Registrar compra", lambda: registrar("Compra"), 20).pack(side=tk.LEFT, padx=6)
botao(frame_btn, "Registrar retirada", lambda: registrar("Retirada"), 20).pack(side=tk.LEFT, padx=6)
botao(frame_btn, "Atualizar tabela", atualizar, 20).pack(side=tk.LEFT, padx=6)

atualizar()
root.mainloop()
